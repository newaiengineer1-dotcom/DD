"""
Document loading + chunking for the data room.

Supported: PDF (pypdf), DOCX (python-docx, incl. tables), XLSX (openpyxl), CSV (summary + head),
TXT / MD / JSON / GEOJSON. Images (PNG/JPG) are listed but not text-indexed (no OCR offline) --
their content is described by captions in the design report; add pytesseract if you need OCR.
"""
from __future__ import annotations
import csv, json, re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterator

from helioassure.config import CHUNK_CHARS, CHUNK_OVERLAP


@dataclass
class Chunk:
    text: str
    source: str            # path relative to the collection root
    folder: str
    locator: str           # "p.3", "sheet:Assumptions", "rows 1-40"
    collection: str = "dataroom"
    meta: dict = field(default_factory=dict)

    @property
    def cite(self) -> str:
        return f"[{self.source} | {self.locator}]"


def _clean(t: str) -> str:
    return re.sub(r"[ \t]+", " ", re.sub(r"\n{3,}", "\n\n", t)).strip()


def load_pdf(path: Path) -> Iterator[tuple[str, str]]:
    from pypdf import PdfReader
    r = PdfReader(str(path))
    for i, pg in enumerate(r.pages, 1):
        try:
            t = pg.extract_text() or ""
        except Exception:
            t = ""
        if t.strip():
            yield f"p.{i}", t


def load_docx(path: Path) -> Iterator[tuple[str, str]]:
    from docx import Document
    d = Document(str(path))
    buf = [p.text for p in d.paragraphs if p.text.strip()]
    yield "body", "\n".join(buf)
    for ti, t in enumerate(d.tables, 1):
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
        yield f"table {ti}", "\n".join(rows)


def load_xlsx(path: Path) -> Iterator[tuple[str, str]]:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), data_only=True, read_only=True)
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            if any(c is not None for c in row):
                rows.append(" | ".join("" if c is None else str(c) for c in row))
        if rows:
            yield f"sheet:{ws.title}", "\n".join(rows)


def load_csv(path: Path) -> Iterator[tuple[str, str]]:
    """Large time-series CSVs are summarised (header, row count, numeric stats, first rows)."""
    with open(path, newline="", encoding="utf-8", errors="ignore") as f:
        rd = list(csv.reader(f))
    if not rd:
        return
    head, body = rd[0], rd[1:]
    out = [f"CSV file with {len(body)} data rows. Columns: {', '.join(head)}"]
    for j, col in enumerate(head):
        vals = []
        for r in body:
            try:
                vals.append(float(r[j]))
            except (ValueError, IndexError):
                pass
        if len(vals) > 5:
            out.append(f"column '{col}': n={len(vals)} min={min(vals):.3f} max={max(vals):.3f} mean={sum(vals)/len(vals):.3f} sum={sum(vals):.1f}")
    out.append("First rows:\n" + "\n".join(", ".join(r) for r in body[:12]))
    yield "summary", "\n".join(out)


def load_text(path: Path) -> Iterator[tuple[str, str]]:
    t = path.read_text(encoding="utf-8", errors="ignore")
    if path.suffix.lower() in (".json", ".geojson"):
        try:
            t = json.dumps(json.loads(t), indent=1)[:20000]
        except Exception:
            pass
    yield "text", t


LOADERS = {".pdf": load_pdf, ".docx": load_docx, ".xlsx": load_xlsx, ".csv": load_csv,
           ".txt": load_text, ".md": load_text, ".json": load_text, ".geojson": load_text}
SKIP_DIRS = {"_figures", "__pycache__", ".git"}


def split(text: str, size: int = CHUNK_CHARS, overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = _clean(text)
    if len(text) <= size:
        return [text] if text else []
    out, i = [], 0
    while i < len(text):
        j = min(len(text), i + size)
        if j < len(text):                       # try to break on paragraph/sentence boundary
            k = max(text.rfind("\n", i + size // 2, j), text.rfind(". ", i + size // 2, j))
            if k > 0:
                j = k + 1
        out.append(text[i:j].strip())
        if j >= len(text):
            break
        i = max(j - overlap, i + 1)
    return [c for c in out if c]


def iter_chunks(root: Path, collection: str = "dataroom") -> Iterator[Chunk]:
    for path in sorted(root.rglob("*")):
        if not path.is_file() or any(p in SKIP_DIRS for p in path.parts):
            continue
        loader = LOADERS.get(path.suffix.lower())
        if not loader:
            continue
        rel = path.relative_to(root).as_posix()
        folder = rel.split("/")[0] if "/" in rel else "."
        try:
            for loc, txt in loader(path):
                for n, piece in enumerate(split(txt)):
                    yield Chunk(piece, rel, folder, loc if n == 0 else f"{loc} (part {n+1})", collection)
        except Exception as e:                  # never let one bad file kill the index
            yield Chunk(f"[UNREADABLE FILE: {e}]", rel, folder, "error", collection)


def inventory(root: Path) -> list[dict]:
    """File-level inventory for the Document Controller agent / UI."""
    rows = []
    for path in sorted(root.rglob("*")):
        if path.is_file() and not any(p in SKIP_DIRS for p in path.parts):
            rel = path.relative_to(root).as_posix()
            rows.append({"path": rel, "folder": rel.split("/")[0] if "/" in rel else ".",
                         "type": path.suffix.lower().lstrip("."), "kb": round(path.stat().st_size / 1024, 1),
                         "indexed_as_text": path.suffix.lower() in LOADERS})
    return rows
